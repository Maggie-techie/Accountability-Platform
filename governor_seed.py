import os
from datetime import datetime, UTC
from openpyxl import load_workbook
from pymongo import MongoClient

# CONFIG
EXCEL_FILE = "Nyeri_Governor_Fiscal_2022_2027(AutoRecovered).xlsx"
DB_NAME = "Accountability"

# CONNECT TO MONGODB
client = MongoClient("mongodb://localhost:27017")
db = client[DB_NAME]
# COLLECTIONS
county_leaders = db["county_leaders"]
county_finances = db["county_finances"]
county_audit = db["county_audit_findings"]
department_absorption = db["department_absorption"]

print("\n--- Seeding Governor Data ---")

# CLEAR OLD DATA
county_leaders.delete_many({})
county_finances.delete_many({})
county_audit.delete_many({})
department_absorption.delete_many({})

# LOAD EXCEL
if not os.path.exists(EXCEL_FILE):
    raise FileNotFoundError(f"{EXCEL_FILE} not found")

wb = load_workbook(EXCEL_FILE, data_only=True)

# ============================
# 1️⃣ GOVERNOR PROFILE
# ============================

county_leaders.insert_one({
    "name": "Mutahi Kahiga",
    "role": "Governor",
    "county": "Nyeri",
    "party": "UDA",
    "tenure": "2022-2027",
    "level": "county",
    "created_at": datetime.now(UTC)
})

print("Governor profile inserted")

# ============================
# 2️⃣ SHEET 1 — FISCAL OVERVIEW
# ============================

sheet1 = wb["Governor & County Overview"]
finance_docs = []

for row in sheet1.iter_rows(min_row=2, values_only=True):

    if not row[0]:
        continue

    # Skip TOTAL rows
    if "TOTAL" in str(row[0]).upper():
        continue

    revenue_source = str(row[0]).strip()
    description = str(row[1]).strip() if row[1] else ""

    years = [
        ("2022/23", row[2]),
        ("2023/24", row[3]),
        ("2024/25", row[4]),
        ("2025/26", row[5])
    ]

    for year, value in years:

        # Skip empty or non-numeric values
        if not isinstance(value, (int, float)):
            continue

        finance_docs.append({
            "county": "Nyeri",
            "governor": "Mutahi Kahiga",
            "revenue_source": revenue_source,
            "description": description,
            "financial_year": year,
            "amount_kshb": float(value),
            "notes": row[7] if len(row) > 7 else "",
            "created_at": datetime.now(UTC)
        })

county_finances.insert_many(finance_docs)
print(f"Inserted {len(finance_docs)} finance records")

# ============================
# 3️⃣ SHEET 2 — AUDIT FINDINGS
# ============================

sheet2 = wb["OAG Audit Findings"]
audit_docs = []

for row in sheet2.iter_rows(min_row=2, values_only=True):

    if not row[0]:
        continue

    try:
        amount = float(row[2]) if isinstance(row[2], (int, float)) else 0
    except:
        amount = 0

    finding_type = "misappropriation"
    if row[5] and str(row[5]).strip():
        finding_type = "correct"

    audit_docs.append({
        "county": "Nyeri",
        "governor": "Mutahi Kahiga",
        "category": str(row[0]).strip(),
        "financial_year": str(row[1]).strip() if row[1] else "",
        "amount_flagged_kshm": amount,
        "severity": str(row[3]).strip() if row[3] else "",
        "finding_type": finding_type,
        "misappropriation_notes": row[4] or "",
        "recommendation": row[5] or "",
        "created_at": datetime.now(UTC)
    })

county_audit.insert_many(audit_docs)
print(f"Inserted {len(audit_docs)} audit records")

# ============================
# 4️⃣ SHEET 3 — DEPARTMENTS
# ============================

sheet3 = wb["Dept Spending Tracker"]
dept_docs = []
current_year = None

for row in sheet3.iter_rows(values_only=True):

    if row[0] and "FY" in str(row[0]):
        current_year = str(row[0]).strip()
        continue

    if not row[0] or "Department" in str(row[0]):
        continue

    if "TOTAL" in str(row[0]).upper():
        continue

    try:
        dept_docs.append({
            "county": "Nyeri",
            "department": str(row[0]).strip(),
            "financial_year": current_year,

            "approved_budget_kshm": float(row[1]) if isinstance(row[1], (int, float)) else 0,
            "q3_spend_kshm": float(row[2]) if isinstance(row[2], (int, float)) else 0,
            "absorption_rate": float(row[3]) if isinstance(row[3], (int, float)) else 0,
            "remaining_budget_kshm": float(row[4]) if isinstance(row[4], (int, float)) else 0,
            "projected_year_end": float(row[5]) if isinstance(row[5], (int, float)) else 0,

            "status": str(row[6]).strip() if row[6] else "",
            "issues": str(row[7]).strip() if row[7] else "",

            "created_at": datetime.now(UTC)
        })

    except Exception as e:
        print(f"Skipping row: {row} → {e}")

department_absorption.insert_many(dept_docs)
print(f"Inserted {len(dept_docs)} department records")

# ============================
# DONE
# ============================

print("\n✅ Governor data successfully seeded!")