import argparse
import os
from datetime import datetime
from openpyxl import load_workbook
from pymongo import MongoClient, ASCENDING, TEXT
from pymongo.errors import ConnectionFailure

#  Config 
EXCEL_FILE = "Nyeri_NGCDF_2022_2027.xlsx"   
DB_NAME    = "Accountability"

#  Argument parser 
parser = argparse.ArgumentParser(description="Seed Nyeri NG-CDF data to MongoDB")
parser.add_argument("--uri", default="mongodb://localhost:27017",
                    help='MongoDB URI e.g. "mongodb+srv://user:pass@cluster.mongodb.net"')
args = parser.parse_args()

# Connect 
print("\n Connecting to MongoDB...")
client = MongoClient(args.uri, serverSelectionTimeoutMS=8000)
try:
    client.admin.command("ping")
    print("Connected successfully\n")
except ConnectionFailure as e:
    print(f"Connection failed: {e}")
    raise

db = client[DB_NAME]

# Drop existing collections for clean reseed
for col in ["constituencies", "allocations", "audit_findings", "mps"]:
    db[col].drop()
print("Cleared existing collections\n")

# STEP 1 — READ EXCEL

if not os.path.exists(EXCEL_FILE):
    print(f"Excel file not found: {EXCEL_FILE}")
    print("Make sure Nyeri_NGCDF_2022_2027.xlsx is in the same folder as this script.")
    raise FileNotFoundError(EXCEL_FILE)
 
print(f"Reading {EXCEL_FILE}...")
wb = load_workbook(EXCEL_FILE, data_only=True)
 
# Sheet 1: Overview — allocations per constituency 
overview_sheet = wb["Overview"]
overview_data  = {}
 
# Row 4-9: A=constituency B=MP C=FY22/23 D=FY23/24 E=FY24/25 F=FY25/26 H=audit status
for row in overview_sheet.iter_rows(min_row=4, max_row=9, values_only=True):
    if not row[0] or "TOTAL" in str(row[0]).upper():
        continue
    name = str(row[0]).strip()
    overview_data[name] = {
        "mp_raw":       str(row[1]).strip() if row[1] else "",
        "FY2022_23":    float(row[2]) if row[2] else 0.0,
        "FY2023_24":    float(row[3]) if row[3] else 0.0,
        "FY2024_25":    float(row[4]) if row[4] else 0.0,
        "FY2025_26":    float(row[5]) if row[5] else 0.0,
        "audit_status": str(row[7]).strip() if row[7] else "Unknown",
    }
 
print(f"  Overview sheet: {len(overview_data)} constituencies")
for name, d in overview_data.items():
    print(f"    {name} | {d['mp_raw']} | FY22/23: {d['FY2022_23']}M")
 
# Sheet 2: Audit Findings Detail 
findings_sheet   = wb["Audit Findings Detail"]
findings_by_slug = {}
 
for row in findings_sheet.iter_rows(min_row=4, max_row=9, values_only=True):
    if not row[0]:
        continue
    name = str(row[0]).strip()
    slug = name.lower().replace(" ", "_")
    findings_by_slug[slug] = {
        "fy_reviewed":           str(row[2]).strip() if row[2] else "",
        "amount_at_risk":        str(row[3]).strip() if row[3] else "",
        "misappropriations_raw": str(row[4]).strip() if row[4] else "",
        "correct_raw":           str(row[5]).strip() if row[5] else "",
    }
 
print(f"\n  Audit findings sheet: {len(findings_by_slug)} records")

# Step 2 - Building Mongo Documents

slug_map = {
    "Tetu":       "tetu",
    "Kieni":      "kieni",
    "Mathira":    "mathira",
    "Othaya":     "othaya",
    "Mukurweini": "mukurweini",
    "Nyeri Town": "nyeri_town",
}
 
opinion_map = {
    "Adverse":      "Adverse",
    "Qualified":    "Qualified",
    "Unqualified*": "Unqualified",
    "Unqualified":  "Unqualified",
}
 
def parse_bullets(raw):
    """Split bullet-point cell text into a clean list."""
    if not raw:
        return []
    lines = [l.strip().lstrip("•").strip() for l in raw.split("\n") if l.strip()]
    return [l for l in lines if l]
 
 

constituency_docs = []
 
for name, ov in overview_data.items():
    slug   = slug_map.get(name, name.lower().replace(" ", "_"))
    detail = findings_by_slug.get(slug, {})
    status = ov["audit_status"].replace("*", "").strip()
 
    misappropriations = parse_bullets(detail.get("misappropriations_raw", ""))
    correct           = parse_bullets(detail.get("correct_raw", ""))
 
    doc = {
        "slug":        slug,
        "name":        name,
        "county":      "Nyeri",
        "county_code": "019",
        "mp": {
            "name":   ov["mp_raw"],
            "party":  "UDA",
            "tenure": "2022-2027",
            "gender": "Male",
        },
        "audit_status": status,
        "oag_opinion":  opinion_map.get(ov["audit_status"], "Qualified"),
        "allocations_ksm": {
            "FY2022_23": ov["FY2022_23"],
            "FY2023_24": ov["FY2023_24"],
            "FY2024_25": ov["FY2024_25"],
            "FY2025_26": ov["FY2025_26"],
        },
        "misappropriations": [
            {
                "id":           f"{slug.upper()}-M-{str(i+1).zfill(3)}",
                "finding":      text,
                "finding_type": "misappropriation",
            }
            for i, text in enumerate(misappropriations)
        ],
        "correct_appropriations": [
            {
                "id":           f"{slug.upper()}-C-{str(i+1).zfill(3)}",
                "note":         text,
                "finding_type": "correct_appropriation",
            }
            for i, text in enumerate(correct)
        ],
        "fy_reviewed":    detail.get("fy_reviewed", ""),
        "amount_at_risk": detail.get("amount_at_risk", ""),
        "data_sources": [
            "Nyeri_NGCDF_2022_2027.xlsx",
            "OAG NG-CDF Constituency Reports",
            "Bunge Library Parliament IR",
        ],
        "created_at": datetime.utcnow(),
        "updated_at": datetime.utcnow(),
    }
    constituency_docs.append(doc)
 
print(f"\nBuilt {len(constituency_docs)} constituency documents")
 
#  Allocations (flat — one per constituency per FY) 
fy_display = {
    "FY2022_23": "2022/23",
    "FY2023_24": "2023/24",
    "FY2024_25": "2024/25",
    "FY2025_26": "2025/26",
}
 
allocation_docs = []
for doc in constituency_docs:
    for fy_key, label in fy_display.items():
        allocation_docs.append({
            "constituency":      doc["name"],
            "constituency_slug": doc["slug"],
            "mp_name":           doc["mp"]["name"],
            "fy_key":            fy_key,
            "fy_display":        label,
            "amount_kshm":        doc["allocations_ksm"][fy_key],
            "county":            "Nyeri",
            "county_code":       "019",
        })
 
print(f"Built {len(allocation_docs)} allocation documents")
 
# Audit Findings (flat — one per finding) 
finding_docs = []
for doc in constituency_docs:
    for m in doc["misappropriations"]:
        finding_docs.append({
            **m,
            "constituency":      doc["name"],
            "constituency_slug": doc["slug"],
            "mp_name":           doc["mp"]["name"],
            "county":            "Nyeri",
        })
    for c in doc["correct_appropriations"]:
        finding_docs.append({
            **c,
            "constituency":      doc["name"],
            "constituency_slug": doc["slug"],
            "mp_name":           doc["mp"]["name"],
            "county":            "Nyeri",
        })
 
print(f"Built {len(finding_docs)} audit finding documents")
 
# MPs (summary scorecard) 
mp_docs = []
for doc in constituency_docs:
    total = sum(doc["allocations_ksm"].values())
    mp_docs.append({
        "name":                        doc["mp"]["name"],
        "party":                       doc["mp"]["party"],
        "tenure":                      doc["mp"]["tenure"],
        "constituency":                doc["name"],
        "constituency_slug":           doc["slug"],
        "county":                      "Nyeri",
        "county_code":                 "019",
        "total_allocation_kshm":        round(total, 2),
        "misappropriation_count":      len(doc["misappropriations"]),
        "correct_appropriation_count": len(doc["correct_appropriations"]),
        "oag_opinion":                 doc["oag_opinion"],
        "audit_status":                doc["audit_status"],
        "amount_at_risk":              doc["amount_at_risk"],
    })
 
print(f"Built {len(mp_docs)} MP documents")
 
# STEP 3 — INSERT INTO MONGODB

print("\nSeeding to MongoDB Atlas...")
 
db["constituencies"].insert_many(constituency_docs)
print(f"  constituencies  — {len(constituency_docs)} inserted")
 
db["allocations"].insert_many(allocation_docs)
print(f"  allocations     — {len(allocation_docs)} inserted")
 
db["audit_findings"].insert_many(finding_docs)
print(f"  audit_findings  — {len(finding_docs)} inserted")
 
db["mps"].insert_many(mp_docs)
print(f"  mps             — {len(mp_docs)} inserted")

# STEP 4 — INDEXES

print("\nCreating indexes...")
 
db["constituencies"].create_index([("slug", ASCENDING)], unique=True)
db["constituencies"].create_index([("audit_status", ASCENDING)])
db["constituencies"].create_index([("name", TEXT), ("mp.name", TEXT)])
db["constituencies"].create_index([("county_code", ASCENDING)])
 
db["allocations"].create_index([("constituency_slug", ASCENDING), ("fy_key", ASCENDING)])
db["allocations"].create_index([("fy_key", ASCENDING)])
db["allocations"].create_index([("county_code", ASCENDING)])
 
db["audit_findings"].create_index([("constituency_slug", ASCENDING)])
db["audit_findings"].create_index([("finding_type", ASCENDING)])
db["audit_findings"].create_index([("mp_name", TEXT), ("finding", TEXT)])
db["audit_findings"].create_index([("county_code", ASCENDING)])
 
db["mps"].create_index([("constituency_slug", ASCENDING)], unique=True)
db["mps"].create_index([("oag_opinion", ASCENDING)])
db["mps"].create_index([("county_code", ASCENDING)])
 
# Create AI cache collection indexes
db["ai_cache"].create_index([("cache_key", ASCENDING)], unique=True)
db["ai_cache"].create_index([("expires_at")], expireAfterSeconds=0)

print("  All indexes created")

# STEP 5 — VERIFY

print(f"\n── '{DB_NAME}' collection counts ────────────────────")
for col in ["constituencies", "allocations", "audit_findings", "mps"]:
    print(f"   {col:<25} {db[col].count_documents({})} documents")
 
print("\nDone. Your database is ready on Compass.")
print(f"Verify in mongosh:")
print(f'  use {DB_NAME}')
print(f'  show collections')
print(f'  db.constituencies.find().pretty()')
 
 
